import pandas as pd
import numpy as np
import os
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl import Workbook
import datetime
import yaml
from utils import should_use_excel, load_cuenca_config

# SCS Type II cumulative mass distribution (25 values, hours 0–24)
_SCS_TYPE_II = np.array([
    0.000, 0.011, 0.022, 0.034, 0.048, 0.063,
    0.080, 0.098, 0.120, 0.147, 0.181, 0.235,
    0.663, 0.772, 0.820, 0.854, 0.880, 0.898,
    0.916, 0.930, 0.944, 0.958, 0.971, 0.983, 1.000
])


def get_cuenca_from_path(path):
    """Extraer nombre de cuenca del path de salida o entrada"""
    parts = path.replace('\\', '/').split('/')
    if 'results' in parts:
        idx = parts.index('results')
        if idx + 1 < len(parts):
            return parts[idx + 1]
    return None

def calcular_tiempo_concentracion(longitud_km, pendiente_porcentaje, cn):
    """Calcula el tiempo de concentración usando la fórmula SCS"""
    L_metros = longitud_km * 1000
    S_porcentaje = pendiente_porcentaje
    tc_horas = (4.3611 * (L_metros**0.8) * ((1000/cn - 9)**0.7)) / (1900 * (S_porcentaje**0.5))
    return tc_horas

def leer_parametros_cuenca(parametros_path, cuenca_config):
    """
    Lee los parámetros geomorfológicos, usando el flag usar_parametros_excel para decidir la fuente
    """
    usar_excel = should_use_excel(parametros_path, cuenca_config)

    if usar_excel:
        try:
            df_parametros = pd.read_excel(parametros_path, index_col=0)
            area_km2 = df_parametros.loc["Área (Km²)", "Valor"]
            longitud_km = df_parametros.loc["Longitud del cauce principal (Km)", "Valor"]
            pendiente_porcentaje = df_parametros.loc["Pendiente media de la cuenca (%)", "Valor"]
            return area_km2, longitud_km, pendiente_porcentaje
        except Exception as e:
            usar_excel = False

    if not usar_excel:
        area_km2 = cuenca_config.get('area_cuenca_km2', 2.5)
        longitud_km = cuenca_config.get('longitud_cauce_principal_km', 3.2)
        pendiente_porcentaje = cuenca_config.get('pendiente_media_cuenca_pct', 8.5)
        return area_km2, longitud_km, pendiente_porcentaje

def leer_hietograma_idf(hietograma_excel_path):
    """Lee el hietograma de precipitación del archivo generado por curvas_IDF"""
    try:
        df_hietograma = pd.read_excel(hietograma_excel_path)

        if 'Tiempo_horas' in df_hietograma.columns and 'Precipitacion_mm' in df_hietograma.columns:
            tiempos = df_hietograma['Tiempo_horas'].values
            precipitacion = df_hietograma['Precipitacion_mm'].values
        elif 'Duracion_min' in df_hietograma.columns and 'Precipitacion_mm' in df_hietograma.columns:
            tiempos = df_hietograma['Duracion_min'].values / 60.0
            precipitacion = df_hietograma['Precipitacion_mm'].values
        elif 'Duracion_min' in df_hietograma.columns and 'Precipitacion_Incremental_mm' in df_hietograma.columns:
            tiempos = df_hietograma['Duracion_min'].values / 60.0
            precipitacion = df_hietograma['Precipitacion_Incremental_mm'].values
        else:
            return None, None

        return tiempos, precipitacion
    except Exception as e:
        return None, None

def calcular_lluvia_efectiva_incremental_scs(precip_mm_inc, cn):
    """Calcula la lluvia efectiva incremental usando el método SCS estándar"""
    S = (25400 / cn) - 254
    Ia = 0.2 * S
    P_acum = 0
    Pe_acum = 0
    lluvia_efectiva = []

    for i, P_inc in enumerate(precip_mm_inc):
        P_acum += P_inc

        if P_acum <= Ia:
            Pe_actual = 0
        else:
            Pe_actual = ((P_acum - Ia)**2) / (P_acum - Ia + S)

        Pe_inc = Pe_actual - Pe_acum
        if Pe_inc < 0:
            Pe_inc = 0
        lluvia_efectiva.append(Pe_inc)
        Pe_acum = Pe_actual

    return np.array(lluvia_efectiva)

def construir_hidrograma_unitario_scs(area_km2, t_conc_horas, dt_horas, duracion_hietograma_horas=None, tiempo_inicio=0.0):
    """Construye el hidrograma unitario SCS según el método estándar"""
    Tp = dt_horas / 2 + 0.6 * t_conc_horas
    Qp_unit = (0.208 * area_km2) / Tp

    if duracion_hietograma_horas is not None:
        duracion = int(np.ceil(duracion_hietograma_horas / dt_horas)) + 1
    else:
        duracion = int(np.ceil(5 * Tp / dt_horas)) + 1

    tiempos = tiempo_inicio + np.arange(duracion) * dt_horas

    caudal = np.zeros_like(tiempos)
    for i, t in enumerate(tiempos):
        t_rel = t - tiempo_inicio

        if t_rel <= 0:
            caudal[i] = 0
        elif t_rel <= Tp:
            caudal[i] = Qp_unit * (t_rel / Tp) * np.exp(1 - t_rel / Tp)
        else:
            caudal[i] = Qp_unit * np.exp(1 - t_rel / Tp)

    return tiempos, caudal

def calcular_t_lag(t_conc_horas):
    return 0.6 * t_conc_horas * 60

def desplazar_hidrograma(df_hid, t_lag_min, dt_horas):
    t_lag_horas = t_lag_min / 60
    n_shift = int(round(t_lag_horas / dt_horas))
    caudal = df_hid['Caudal_Diseño_m3s'].values
    caudal_shifted = np.concatenate((np.zeros(n_shift), caudal))
    tiempos_shifted = np.arange(len(caudal_shifted)) * dt_horas
    return pd.DataFrame({'Tiempo_horas': tiempos_shifted, 'Caudal_Diseño_m3s': caudal_shifted})


# ---------------------------------------------------------------------------
# Gumbel + SCS Type II path (for cuenca_ya_delimitada=true)
# ---------------------------------------------------------------------------

def calcular_precipitacion_gumbel(precip_path, periodos=(50, 100, 200)):
    """Gumbel EV-I frequency analysis on annual max 24h precipitation series."""
    df = pd.read_excel(precip_path)
    col_p = df.columns[-1]  # WaterKu 2-col format: (Año, Precipitación)
    serie = pd.to_numeric(df[col_p], errors='coerce').dropna().values
    media = serie.mean()
    std = serie.std(ddof=1)
    alpha = np.pi / (std * np.sqrt(6))
    mu = media - 0.5772 / alpha
    resultados = {}
    for T in periodos:
        yT = -np.log(-np.log(1.0 - 1.0 / T))
        resultados[T] = float(mu + yT / alpha)
    return resultados, dict(media=float(media), std=float(std), alpha=float(alpha), mu=float(mu), n=int(len(serie)))


def generar_hietograma_scs_type2(P24h, dt_horas=1.0):
    """24-hour incremental precipitation using SCS Type II distribution."""
    P_inc = np.diff(_SCS_TYPE_II * P24h)  # 24 incremental values
    tiempos = np.arange(1, 25) * dt_horas
    return pd.DataFrame({'Tiempo_horas': tiempos, 'Precipitacion_mm': P_inc})


def _convolve_hydrograph(area_km2, t_conc, cn, precipitacion, dt_horas):
    """Convolve effective rainfall with SCS UH. Returns df_hid."""
    lluvia_efectiva = calcular_lluvia_efectiva_incremental_scs(precipitacion, cn)
    duracion_hietograma = len(precipitacion) * dt_horas
    tiempos_hu, caudales_hu = construir_hidrograma_unitario_scs(
        area_km2, t_conc, dt_horas, duracion_hietograma, 0.0
    )
    hidrograma_full = np.convolve(lluvia_efectiva, caudales_hu)
    Tp = dt_horas / 2 + 0.6 * t_conc
    duracion_total = 24 + int(np.ceil(3 * Tp))
    max_puntos = min(duracion_total, len(hidrograma_full))
    hidrograma = hidrograma_full[:max_puntos]
    tiempos_hid = np.arange(len(hidrograma)) * dt_horas
    df_hid = pd.DataFrame({'Tiempo_horas': tiempos_hid, 'Caudal_Diseño_m3s': hidrograma})
    return df_hid, lluvia_efectiva


def procesar_cuenca_gumbel(area_km2, longitud_km, pendiente_pct, cn, precip_anual_path,
                            periodos=(50, 100, 200), dt_horas=1.0, output_dir="."):
    """Run SCS hydrograph for multiple return periods using Gumbel EV-I + SCS Type II."""
    os.makedirs(output_dir, exist_ok=True)

    p24h_dict, gumbel_stats = calcular_precipitacion_gumbel(precip_anual_path, periodos)
    t_conc = calcular_tiempo_concentracion(longitud_km, pendiente_pct, cn)

    resultados_tr = {}
    for Tr in periodos:
        P24h = p24h_dict[Tr]
        df_hiet = generar_hietograma_scs_type2(P24h, dt_horas)
        precipitacion = df_hiet['Precipitacion_mm'].values
        df_hid, lluvia_efectiva = _convolve_hydrograph(area_km2, t_conc, cn, precipitacion, dt_horas)
        resultados_tr[Tr] = dict(
            pico=float(df_hid['Caudal_Diseño_m3s'].max()),
            P24h=P24h,
            df_hid=df_hid,
            precipitacion=precipitacion,
            lluvia_efectiva=lluvia_efectiva,
        )

    archivo_excel = os.path.join(output_dir, "analisis_hidrograma_scs_cuenca.xlsx")
    _escribir_excel_gumbel(archivo_excel, resultados_tr, gumbel_stats, area_km2, t_conc, cn)

    return {Tr: resultados_tr[Tr]['pico'] for Tr in periodos}


def _escribir_excel_gumbel(archivo_excel, resultados_tr, gumbel_stats, area_km2, t_conc, cn):
    """Write multi-Tr Excel: Resumen sheet + one sheet per return period."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.chart import LineChart, Reference

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Resumen Caudales"

    # --- Resumen sheet ---
    headers_sum = ["Período de Retorno (años)", "P24h Gumbel (mm)", "Caudal Pico (m³/s)"]
    for col, h in enumerate(headers_sum, 1):
        c = ws_sum.cell(row=1, column=col, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")

    for row_idx, Tr in enumerate(sorted(resultados_tr.keys()), 2):
        res = resultados_tr[Tr]
        ws_sum.cell(row=row_idx, column=1, value=Tr)
        ws_sum.cell(row=row_idx, column=2, value=round(res['P24h'], 2))
        ws_sum.cell(row=row_idx, column=3, value=round(res['pico'], 3))

    # Gumbel stats
    row_stats = len(resultados_tr) + 3
    ws_sum.cell(row=row_stats, column=1, value="Estadísticos Gumbel").font = Font(bold=True)
    stats_rows = [
        ("N datos", gumbel_stats['n']),
        ("Media (mm)", round(gumbel_stats['media'], 3)),
        ("Desv. Est. (mm)", round(gumbel_stats['std'], 3)),
        ("Alpha", round(gumbel_stats['alpha'], 6)),
        ("Mu", round(gumbel_stats['mu'], 3)),
        ("Tc (horas)", round(t_conc, 3)),
        ("Área (km²)", area_km2),
        ("CN", cn),
    ]
    for i, (label, val) in enumerate(stats_rows):
        ws_sum.cell(row=row_stats + 1 + i, column=1, value=label)
        ws_sum.cell(row=row_stats + 1 + i, column=2, value=val)

    ws_sum.column_dimensions['A'].width = 28
    ws_sum.column_dimensions['B'].width = 20
    ws_sum.column_dimensions['C'].width = 22

    # --- One sheet per Tr ---
    fecha_base = datetime.datetime(2000, 1, 1)
    for Tr in sorted(resultados_tr.keys()):
        res = resultados_tr[Tr]
        df_hid = res['df_hid']
        precipitacion = res['precipitacion']
        lluvia_efectiva = res['lluvia_efectiva']

        ws = wb.create_sheet(f"Tr{Tr:03d}")
        headers = ['Hora', 'Precipitación (mm)', 'Pérdidas (mm)', 'Exceso (mm)', 'Caudal (m³/s)']
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=col, value=h)
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center")

        for i in range(len(df_hid)):
            t = df_hid['Tiempo_horas'].iloc[i]
            fecha_hora = fecha_base + datetime.timedelta(hours=t)
            idx_p = int(t) - 1  # hietogram at hours 1-24
            p_inc = float(precipitacion[idx_p]) if 0 <= idx_p < len(precipitacion) else 0.0
            ll_eff = float(lluvia_efectiva[idx_p]) if 0 <= idx_p < len(lluvia_efectiva) else 0.0
            caudal = df_hid['Caudal_Diseño_m3s'].iloc[i]

            ws.cell(row=i+2, column=1, value=fecha_hora.strftime('%H:%M'))
            ws.cell(row=i+2, column=2, value=round(p_inc, 2))
            ws.cell(row=i+2, column=3, value=round(p_inc - ll_eff, 2))
            ws.cell(row=i+2, column=4, value=round(ll_eff, 2))
            ws.cell(row=i+2, column=5, value=round(caudal, 3))

        # chart
        ws_chart = wb.create_sheet(f"Graf_Tr{Tr:03d}")
        ws_chart.cell(row=1, column=1, value="Tiempo (h)").font = Font(bold=True)
        ws_chart.cell(row=1, column=2, value=f"Q Tr{Tr} (m³/s)").font = Font(bold=True)
        n_pts = len(df_hid)
        for i in range(n_pts):
            ws_chart.cell(row=i+2, column=1, value=df_hid['Tiempo_horas'].iloc[i])
            ws_chart.cell(row=i+2, column=2, value=df_hid['Caudal_Diseño_m3s'].iloc[i])

        chart = LineChart()
        chart.title = f"Hidrograma Tr={Tr} años"
        chart.style = 2
        chart.y_axis.title = 'Caudal (m³/s)'
        chart.x_axis.title = 'Tiempo (horas)'
        chart.legend = None
        data_ref = Reference(ws_chart, min_col=2, min_row=1, max_row=n_pts+1)
        chart.add_data(data_ref, titles_from_data=True)
        cats = Reference(ws_chart, min_col=1, min_row=2, max_row=n_pts+1)
        chart.set_categories(cats)
        if chart.series:
            chart.series[0].graphicalProperties.line.solidFill = "0070C0"
            chart.series[0].graphicalProperties.line.width = 25000
        chart.height = 10
        chart.width = 15
        ws_chart.add_chart(chart, "D2")

    wb.save(archivo_excel)


# ---------------------------------------------------------------------------
# IDF / Bloque Alterno path (for cuenca_ya_delimitada=false)
# ---------------------------------------------------------------------------

def procesar_cuenca_integrada(parametros_excel_path, hietograma_excel_path, cuenca_config, output_dir, dt_horas=1.0):
    """Procesa una cuenca usando datos integrados de otros scripts"""
    os.makedirs(output_dir, exist_ok=True)

    cn = cuenca_config.get('cn_default', 85)
    area_km2, longitud_km, pendiente_porcentaje = leer_parametros_cuenca(parametros_excel_path, cuenca_config)
    if area_km2 is None:
        raise ValueError("No se pudieron leer los parámetros de la cuenca")

    t_conc = calcular_tiempo_concentracion(longitud_km, pendiente_porcentaje, cn)

    tiempos, precipitacion = leer_hietograma_idf(hietograma_excel_path)
    if tiempos is None:
        raise ValueError("No se pudo leer el hietograma de precipitación")

    tiempo_inicio = 0.0
    lluvia_efectiva = calcular_lluvia_efectiva_incremental_scs(precipitacion, cn)
    tiempos_precipitacion = np.arange(tiempo_inicio, tiempo_inicio + len(precipitacion) * dt_horas, dt_horas)
    duracion_hietograma = len(precipitacion) * dt_horas
    tiempos_hu, caudales_hu = construir_hidrograma_unitario_scs(area_km2, t_conc, dt_horas, duracion_hietograma, tiempo_inicio)

    hidrograma_full = np.convolve(lluvia_efectiva, caudales_hu)

    duracion_lluvia = 24
    Tp = dt_horas / 2 + 0.6 * t_conc
    duracion_total = duracion_lluvia + int(np.ceil(3 * Tp))
    max_puntos = min(duracion_total, len(hidrograma_full))

    hidrograma = hidrograma_full[:max_puntos]
    tiempos_hid = tiempo_inicio + np.arange(len(hidrograma)) * dt_horas
    df_hidrograma = pd.DataFrame({'Tiempo_horas': tiempos_hid, 'Caudal_Diseño_m3s': hidrograma})

    pico = df_hidrograma['Caudal_Diseño_m3s'].max()
    tiempo_pico = df_hidrograma.loc[df_hidrograma['Caudal_Diseño_m3s'].idxmax(), 'Tiempo_horas']

    nombre_cuenca = "cuenca"
    archivo_excel = os.path.join(output_dir, f"analisis_hidrograma_scs_{nombre_cuenca}.xlsx")

    fecha_base = datetime.datetime(2000, 1, 1)

    tabla_principal = []

    for i in range(len(df_hidrograma)):
        tiempo_actual = df_hidrograma['Tiempo_horas'].iloc[i]
        fecha_hora = fecha_base + datetime.timedelta(hours=tiempo_actual)

        indice_precip = int(tiempo_actual - tiempo_inicio)

        if 0 <= indice_precip < len(precipitacion):
            precip_inc = precipitacion[indice_precip]
            lluvia_eff = lluvia_efectiva[indice_precip]
        else:
            precip_inc = 0.0
            lluvia_eff = 0.0

        caudal = df_hidrograma['Caudal_Diseño_m3s'].iloc[i]

        tabla_principal.append({
            'Fecha': fecha_hora.strftime('%d%b.%Y'),
            'Hora': fecha_hora.strftime('%H:%M'),
            'Precipitación (MM)': round(precip_inc, 2),
            'Pérdidas (MM)': round(precip_inc - lluvia_eff, 2),
            'Exceso (MM)': round(lluvia_eff, 2),
            'Caudal Directo (M3/S)': round(caudal, 1)
        })

    df_tabla = pd.DataFrame(tabla_principal)

    volumen_precipitacion = precipitacion.sum() * area_km2 * 1000
    volumen_perdidas = (precipitacion.sum() - lluvia_efectiva.sum()) * area_km2 * 1000
    volumen_exceso = lluvia_efectiva.sum() * area_km2 * 1000
    volumen_descarga = np.trapezoid(df_hidrograma['Caudal_Diseño_m3s'],
                                   df_hidrograma['Tiempo_horas']) * 3600

    resumen_computado = {
        'Parámetro': [
            'Caudal Pico:',
            'Volumen Precipitación:',
            'Volumen Pérdidas:',
            'Volumen Exceso:',
            'Volumen Descarga:'
        ],
        'Valor': [
            f'{pico:.1f} (M3/S)',
            f'{volumen_precipitacion/1000:.1f} (1000 M3)',
            f'{volumen_perdidas/1000:.1f} (1000 M3)',
            f'{volumen_exceso/1000:.1f} (1000 M3)',
            f'{volumen_descarga/1000:.1f} (1000 M3)'
        ]
    }

    df_resumen = pd.DataFrame(resumen_computado)

    crear_excel_completo(archivo_excel, df_tabla, df_resumen, df_hidrograma, precipitacion, df_hidrograma['Tiempo_horas'])

    return pico, tiempo_pico

def crear_excel_completo(archivo_excel, df_tabla, df_resumen, df_hidrograma, precipitacion, tiempos):
    """Crea un archivo Excel completo con tabla de datos, resumen y gráfico"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill
        from openpyxl.chart import LineChart, Reference
    except ImportError:
        with pd.ExcelWriter(archivo_excel, engine='xlsxwriter') as writer:
            df_tabla.to_excel(writer, sheet_name='Datos Hidrograma', index=False)
            df_resumen.to_excel(writer, sheet_name='Resultados Computados', index=False)
        return

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Datos Hidrograma"

    headers = ['Fecha', 'Hora', 'Precipitación (MM)', 'Pérdidas (MM)', 'Exceso (MM)', 'Caudal Directo (M3/S)']
    for col, header in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.font = Font(color="000000", bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row, (_, values) in enumerate(df_tabla.iterrows(), 2):
        for col, value in enumerate(values, 1):
            cell = ws1.cell(row=row, column=col, value=value)
            if col > 2:
                cell.alignment = Alignment(horizontal="right")
                if isinstance(value, (int, float)):
                    if col == 6:
                        cell.number_format = '0.0'
                    else:
                        cell.number_format = '0.00'
            else:
                cell.alignment = Alignment(horizontal="center")

    column_widths = [12, 8, 18, 15, 13, 20]
    for col, width in enumerate(column_widths, 1):
        column_letter = chr(64 + col)
        ws1.column_dimensions[column_letter].width = width

    ws2 = wb.create_sheet("Resultados")

    ws2.cell(row=1, column=1, value="Resultados").font = Font(size=11, bold=True)
    ws2.merge_cells('A1:B1')
    ws2.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    for row, (_, values) in enumerate(df_resumen.iterrows(), 3):
        for col, value in enumerate(values, 1):
            if pd.notna(value) and value != '':
                cell = ws2.cell(row=row, column=col, value=value)
                if col == 1:
                    cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="left")

    ws2.column_dimensions['A'].width = 25
    ws2.column_dimensions['B'].width = 25

    ws3 = wb.create_sheet("Gráfico Hidrograma")

    ws3.cell(row=1, column=1, value="Tiempo (horas)").font = Font(bold=True)
    ws3.cell(row=1, column=2, value="Caudal (m³/s)").font = Font(bold=True)

    max_puntos = len(df_hidrograma)
    for i in range(max_puntos):
        ws3.cell(row=i+2, column=1, value=df_hidrograma['Tiempo_horas'].iloc[i])
        ws3.cell(row=i+2, column=2, value=df_hidrograma['Caudal_Diseño_m3s'].iloc[i])

    chart = LineChart()
    chart.title = "Hidrograma de Caudales"
    chart.style = 2
    chart.y_axis.title = 'Caudal (m³/s)'
    chart.x_axis.title = 'Tiempo (horas)'
    chart.legend = None

    data_caudal = Reference(ws3, min_col=2, min_row=1, max_row=max_puntos+1)
    chart.add_data(data_caudal, titles_from_data=True)

    cats = Reference(ws3, min_col=1, min_row=2, max_row=max_puntos+1)
    chart.set_categories(cats)

    if chart.series:
        s1 = chart.series[0]
        s1.graphicalProperties.line.solidFill = "0000FF"
        s1.graphicalProperties.line.width = 25000

    chart.height = 10
    chart.width = 15
    ws3.add_chart(chart, "D2")

    wb.save(archivo_excel)

def main():
    if 'snakemake' in globals():
        cuenca_id = getattr(snakemake.wildcards, 'cuenca', None) or get_cuenca_from_path(str(snakemake.output[0]))
        cuenca_config, global_config = load_cuenca_config(cuenca_id)
        parametros_excel_path = str(snakemake.input.parametros)
        output_dir = os.path.dirname(str(snakemake.output[0]))

        if cuenca_config.get('cuenca_ya_delimitada', False):
            area_km2, longitud_km, pendiente_pct = leer_parametros_cuenca(parametros_excel_path, cuenca_config)
            cn = cuenca_config.get('cn_default', 85)
            precip_anual_path = str(snakemake.input.precipitacion_anual)
            periodos = tuple(getattr(snakemake.params, 'periodos_retorno', [50, 100]))
            procesar_cuenca_gumbel(
                area_km2, longitud_km, pendiente_pct, cn,
                precip_anual_path, periodos=periodos,
                dt_horas=1.0, output_dir=output_dir
            )
        else:
            hietograma_excel_path = str(snakemake.input.hietograma)
            procesar_cuenca_integrada(parametros_excel_path, hietograma_excel_path, cuenca_config, output_dir, dt_horas=1.0)
    else:
        # Para ejecución directa: usar configuración legacy
        with open('config/config.yaml', 'r', encoding='utf-8') as file:
            global_config = yaml.safe_load(file)
        cuenca_config = global_config  # fallback para compatibilidad
        parametros_excel_path = global_config['parametros_excel']
        hietograma_excel_path = global_config['hietograma_excel']
        output_dir = global_config['hidrograma_scs_output_dir']
        procesar_cuenca_integrada(parametros_excel_path, hietograma_excel_path, cuenca_config, output_dir, dt_horas=1.0)

if __name__ == "__main__":
    main()
